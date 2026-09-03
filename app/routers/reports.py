from fastapi import APIRouter, Depends, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from app.database import get_db
from app.models import Farmer, Booking, Commitment, SeedIssue, SeedPayment, BardanaIssue, Dispatch, DispatchLine, BookingVariety
from app.auth import current_user
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import io

from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.lib.units import cm

router = APIRouter()

def make_xlsx(title, headers, rows, totals=None, filters_applied=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]
    
    ws.cell(1, 1, 'Green Fay Farm Foods').font = Font(bold=True, size=14, color='154D37')
    ws.cell(2, 1, title).font = Font(bold=True, size=12)
    ws.cell(3, 1, f'Generated: {datetime.now().strftime("%d %b %Y %H:%M")}').font = Font(italic=True, color='6F7F78')
    if filters_applied:
        ws.cell(4, 1, f'Filters: {filters_applied}').font = Font(italic=True, color='6F7F78')
    
    header_row = 6
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(header_row, ci, h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='154D37')
        cell.alignment = Alignment(horizontal='center')
    
    ws.freeze_panes = f'A{header_row+1}'
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{header_row}'
    
    for ri, row in enumerate(rows, header_row + 1):
        for ci, val in enumerate(row, 1):
            ws.cell(ri, ci, val)
    
    if totals:
        total_row = header_row + len(rows) + 2
        ws.cell(total_row, 1, 'TOTAL').font = Font(bold=True)
        for ci, val in enumerate(totals, 1):
            if val is not None:
                c = ws.cell(total_row, ci, val)
                c.font = Font(bold=True)
    
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)
    
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def make_pdf(title, headers, rows, totals=None, filters_applied=None, landscape_mode=False):
    buf = io.BytesIO()
    pagesize = landscape(A4) if landscape_mode else A4
    doc = SimpleDocTemplate(buf, pagesize=pagesize, 
                           leftMargin=1.5*cm, rightMargin=1.5*cm,
                           topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    story = []
    
    story.append(Paragraph('Green Fay Farm Foods', styles['Heading1']))
    story.append(Paragraph(title, styles['Heading2']))
    story.append(Paragraph(f'Generated: {datetime.now().strftime("%d %b %Y %H:%M")}', styles['Normal']))
    if filters_applied:
        story.append(Paragraph(f'Filters: {filters_applied}', styles['Normal']))
    story.append(Spacer(1, 0.5*cm))
    
    table_data = [headers] + [list(r) for r in rows]
    if totals:
        table_data.append(['TOTAL'] + list(totals[1:]))
    
    col_count = len(headers)
    col_width = (pagesize[0] - 3*cm) / col_count
    t = Table(table_data, colWidths=[col_width]*col_count, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#154D37')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F5F3EB')]),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#DEDDD3')),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
    ]))
    story.append(t)
    
    doc.build(story)
    buf.seek(0)
    return buf

@router.get('/farmers.xlsx')
def report_farmers_xlsx(village: str = '', status: str = '', db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(Farmer)
    if village: q = q.filter(Farmer.village.ilike(f'%{village}%'))
    if status == 'Active': q = q.filter(Farmer.active == True)
    elif status == 'Inactive': q = q.filter(Farmer.active == False)
    
    headers = ['Farmer Code', 'Name', 'Father/Husband', 'Mobile', 'Village', 'Address', 'Status']
    rows = [[f.farmer_code, f.name, f.relation_name, f.mobile, f.village, f.address, 'Active' if f.active else 'Inactive'] for f in q.order_by(Farmer.name).all()]
    
    buf = make_xlsx('Farmer Register', headers, rows, filters_applied=f'Village={village}')
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=farmers.xlsx'})

@router.get('/bookings.xlsx')
def report_bookings_xlsx(season: str = '', status: str = '', db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(Booking)
    if season: q = q.filter(Booking.season == season)
    if status: q = q.filter(Booking.status == status)
    
    headers = ['Booking Code', 'Farmer', 'Village', 'Season', 'Date', 'Agreement No', 'Acres', 'Status']
    rows = [[b.booking_code, b.farmer.name if b.farmer else '', b.farmer.village if b.farmer else '', b.season, str(b.booking_date), b.agreement_no, b.total_acres, b.status] for b in q.order_by(Booking.id.desc()).all()]
    
    buf = make_xlsx('Booking Register', headers, rows)
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=bookings.xlsx'})

@router.get('/contracts.xlsx')
def report_contracts_xlsx(season: str = '', month: str = '', db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(Commitment).join(Commitment.booking)
    if season: q = q.filter(Booking.season == season)
    if month: q = q.filter(Commitment.contract_month == month)
    
    headers = ['Farmer', 'Village', 'Booking', 'Month', 'Year', 'Variety', 'Contracted Bags', 'Delivered Bags', 'Balance', 'Destination', 'Status']
    rows = []
    tot_contracted = 0
    tot_delivered = 0
    for c in q.all():
        b = c.booking
        variety = c.booking_variety.variety if c.booking_variety else ''
        received = db.query(func.sum(DispatchLine.bags)).join(Dispatch).filter(DispatchLine.commitment_id == c.id, Dispatch.status == 'Finalized').scalar() or 0
        bal = c.contracted_bags - received
        status = 'Completed' if bal <= 0 else ('In Progress' if received > 0 else 'Not Started')
        tot_contracted += c.contracted_bags
        tot_delivered += received
        rows.append([b.farmer.name if b.farmer else '', b.farmer.village if b.farmer else '', b.booking_code, c.contract_month, c.contract_year, variety, c.contracted_bags, received, bal, c.destination, status])
    
    buf = make_xlsx('Month-wise Contracts', headers, rows, totals=['', '', '', '', '', '', tot_contracted, tot_delivered, tot_contracted - tot_delivered, '', ''])
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=contracts.xlsx'})

@router.get('/seed-distribution.xlsx')
def report_seed_dist_xlsx(db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(SeedIssue)
    headers = ['Farmer', 'Booking', 'Variety', 'Date', 'Packets', 'Value', 'Status']
    rows = [[db.get(Farmer, s.farmer_id).name if s.farmer_id else '', s.booking_id, s.variety, str(s.issue_date), s.packets, s.total_value, s.status] for s in q.all()]
    
    buf = make_xlsx('Seed Distribution', headers, rows)
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=seed_distribution.xlsx'})

@router.get('/seed-outstanding.xlsx')
def report_seed_out_xlsx(db: Session = Depends(get_db), user = Depends(current_user)):
    farmers = db.query(Farmer).filter(Farmer.active == True).all()
    headers = ['Farmer Code', 'Name', 'Village', 'Total Seed Value', 'Paid', 'Outstanding']
    rows = []
    tot_val = 0; tot_paid = 0; tot_out = 0
    for f in farmers:
        issues = db.query(SeedIssue).filter(SeedIssue.farmer_id == f.id, SeedIssue.status == 'Active').all()
        val = sum(i.total_value for i in issues)
        paid = db.query(func.sum(SeedPayment.amount)).filter(SeedPayment.farmer_id == f.id, SeedPayment.status == 'Active').scalar() or 0
        bal = val - paid
        if val > 0 or paid > 0:
            rows.append([f.farmer_code, f.name, f.village, val, paid, bal])
            tot_val += val; tot_paid += paid; tot_out += bal
    
    buf = make_xlsx('Seed Outstanding', headers, rows, totals=['', '', '', tot_val, tot_paid, tot_out])
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=seed_outstanding.xlsx'})

@router.get('/seed-outstanding.pdf')
def report_seed_out_pdf(db: Session = Depends(get_db), user = Depends(current_user)):
    farmers = db.query(Farmer).filter(Farmer.active == True).all()
    headers = ['Farmer Code', 'Name', 'Village', 'Total Seed Value', 'Paid', 'Outstanding']
    rows = []
    tot_val = 0; tot_paid = 0; tot_out = 0
    for f in farmers:
        issues = db.query(SeedIssue).filter(SeedIssue.farmer_id == f.id, SeedIssue.status == 'Active').all()
        val = sum(i.total_value for i in issues)
        paid = db.query(func.sum(SeedPayment.amount)).filter(SeedPayment.farmer_id == f.id, SeedPayment.status == 'Active').scalar() or 0
        bal = val - paid
        if val > 0 or paid > 0:
            rows.append([f.farmer_code, f.name, f.village, val, paid, bal])
            tot_val += val; tot_paid += paid; tot_out += bal
            
    buf = make_pdf('Seed Outstanding Report', headers, rows, totals=['', '', '', tot_val, tot_paid, tot_out])
    return StreamingResponse(buf, media_type='application/pdf', headers={'Content-Disposition': 'attachment; filename=seed_outstanding.pdf'})

@router.get('/bardana.xlsx')
def report_bardana_xlsx(db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(BardanaIssue)
    headers = ['Farmer', 'Booking', 'Date', 'Month', 'Bags Issued', 'Status']
    rows = [[db.get(Farmer, b.farmer_id).name if b.farmer_id else '', b.booking_id, str(b.issue_date), b.contract_month, b.bags_issued, b.status] for b in q.all()]
    
    buf = make_xlsx('Bardana Issued', headers, rows)
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=bardana.xlsx'})

@router.get('/procurement-due.xlsx')
def report_proc_due_xlsx(month: str = '', db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(Commitment)
    if month: q = q.filter(Commitment.contract_month == month)
    headers = ['Farmer', 'Village', 'Booking', 'Month', 'Variety', 'Contracted Bags', 'Received', 'Pending']
    rows = []
    t_c = 0; t_r = 0; t_p = 0
    for c in q.all():
        b = c.booking
        if not b or not b.farmer: continue
        variety = c.booking_variety.variety if c.booking_variety else ''
        received = db.query(func.sum(DispatchLine.bags)).join(Dispatch).filter(DispatchLine.commitment_id == c.id, Dispatch.status == 'Finalized').scalar() or 0
        pending = c.contracted_bags - received
        rows.append([b.farmer.name, b.farmer.village, b.booking_code, c.contract_month, variety, c.contracted_bags, received, pending])
        t_c += c.contracted_bags; t_r += received; t_p += pending
        
    buf = make_xlsx('Procurement Due', headers, rows, totals=['', '', '', '', '', t_c, t_r, t_p])
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=procurement_due.xlsx'})

@router.get('/procurement-due.pdf')
def report_proc_due_pdf(month: str = '', db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(Commitment)
    if month: q = q.filter(Commitment.contract_month == month)
    headers = ['Farmer', 'Village', 'Booking', 'Month', 'Variety', 'Contracted Bags', 'Received', 'Pending']
    rows = []
    t_c = 0; t_r = 0; t_p = 0
    for c in q.all():
        b = c.booking
        if not b or not b.farmer: continue
        variety = c.booking_variety.variety if c.booking_variety else ''
        received = db.query(func.sum(DispatchLine.bags)).join(Dispatch).filter(DispatchLine.commitment_id == c.id, Dispatch.status == 'Finalized').scalar() or 0
        pending = c.contracted_bags - received
        rows.append([b.farmer.name, b.farmer.village, b.booking_code, c.contract_month, variety, c.contracted_bags, received, pending])
        t_c += c.contracted_bags; t_r += received; t_p += pending
        
    buf = make_pdf('Procurement Due', headers, rows, totals=['', '', '', '', '', t_c, t_r, t_p], landscape_mode=True)
    return StreamingResponse(buf, media_type='application/pdf', headers={'Content-Disposition': 'attachment; filename=procurement_due.pdf'})

@router.get('/dispatch-register.xlsx')
def report_dispatch_xlsx(db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(Dispatch)
    headers = ['Code', 'Date', 'Truck No', 'Destination', 'Actual Wt', 'Mandi Wt', 'Status']
    rows = [[d.dispatch_code, str(d.dispatch_date), d.truck_no, d.destination, d.actual_weight_kg, d.mandi_weight_kg, d.status] for d in q.order_by(Dispatch.id.desc()).all()]
    
    buf = make_xlsx('Dispatch Register', headers, rows)
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=dispatch_register.xlsx'})

@router.get('/dispatch-report.pdf')
@router.get('/dispatch-register.pdf')
def report_dispatch_pdf(db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(Dispatch)
    headers = ['Code', 'Date', 'Truck No', 'Destination', 'Actual Wt', 'Mandi Wt', 'Status']
    rows = [[d.dispatch_code, str(d.dispatch_date), d.truck_no, d.destination, d.actual_weight_kg, d.mandi_weight_kg, d.status] for d in q.order_by(Dispatch.id.desc()).all()]
    
    buf = make_pdf('Dispatch Report', headers, rows, landscape_mode=True)
    return StreamingResponse(buf, media_type='application/pdf', headers={'Content-Disposition': 'attachment; filename=dispatch_report.pdf'})

@router.get('/mandi-tax.xlsx')
def report_mandi_tax_xlsx(db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(Dispatch)
    headers = ['Code', 'Date', 'Mandi Wt', 'Mandi Tax', 'Vikas Sulk', 'Status']
    rows = [[d.dispatch_code, str(d.dispatch_date), d.mandi_weight_kg, d.mandi_tax, d.vikas_sulk, d.status] for d in q.order_by(Dispatch.id.desc()).all()]
    
    buf = make_xlsx('Mandi Tax & Vikas Sulk Report', headers, rows)
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=mandi_tax.xlsx'})

@router.get('/mandi-tax.pdf')
def report_mandi_tax_pdf(db: Session = Depends(get_db), user = Depends(current_user)):
    q = db.query(Dispatch)
    headers = ['Code', 'Date', 'Mandi Wt', 'Mandi Tax', 'Vikas Sulk', 'Status']
    rows = [[d.dispatch_code, str(d.dispatch_date), d.mandi_weight_kg, d.mandi_tax, d.vikas_sulk, d.status] for d in q.order_by(Dispatch.id.desc()).all()]
    
    buf = make_pdf('Mandi Tax & Vikas Sulk Report', headers, rows)
    return StreamingResponse(buf, media_type='application/pdf', headers={'Content-Disposition': 'attachment; filename=mandi_tax.pdf'})
